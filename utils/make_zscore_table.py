"""
Z-score table (BtH paper Table 1 style): original BtH models + new 4096-feature
experiments. Each model -> combined score = mean(global, per-target) for MCC &
BEDROC on Random and Temporal splits. Z computed jointly across ALL models per
column (sums to zero). Average = mean of the 4 column z-scores; SEM = std/sqrt(4).

Emits a LaTeX table (best value per column in italics) and prints it for review.
"""
import numpy as np
import openpyxl

XLSX = "output/results_randomnets_beyond_the_hype.xlsx"
COLS = ["MCC random", "BEDROC random", "MCC temporal", "BEDROC temporal"]

### ORIGINAL BTH MODELS ###
OLD_ORDER = ["NB 10 uM", "NB", "RF", "RF_PCM", "SVM", "LR", "DNN", "DNN_MC", "DNN_PCM"]

def load_old():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Combined Z (BtH + new)"]
    vals = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] in OLD_ORDER and isinstance(row[1], (int, float)):
            vals[row[0]] = [float(row[1]), float(row[2]), float(row[3]), float(row[4])]
    return vals

def load_new():
    """New 4096 models: combine global+per-target from Sheet1 (4096 blocks)."""
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    algos = {"RandomForest": "RF (4096)", "DNN": "DNN (4096)",
             "Randomnets (1 layer)": "RN (1 layer)", "Randomnets (2 layers)": "RN (2 layers)"}
    # dataset -> split -> fp -> algo -> (mcc_g, mcc_pt, bed_g, bed_pt)
    data = {}
    dataset = split = fp = None
    for r in rows:
        f = r[0]
        if f in ("Original BTH", "CHEMBL35"):
            dataset = f; data[dataset] = {}
        elif f in ("Random", "Temporal"):
            split = f; data[dataset][split] = {}
        elif f in (256, 4096):
            fp = f; data[dataset][split][fp] = {}
        elif f in algos and None not in (r[2], r[3], r[4], r[5]) and r[2] != "N/A":
            data[dataset][split][fp][f] = (float(r[2]), float(r[3]), float(r[4]), float(r[5]))
    d = data["Original BTH"]                      # BtH benchmark dataset
    out = {}
    for a, label in algos.items():
        rnd = d["Random"][4096].get(a)
        tmp = d["Temporal"][4096].get(a)
        if rnd is None or tmp is None:
            continue
        out[label] = [(rnd[0] + rnd[1]) / 2, (rnd[2] + rnd[3]) / 2,
                      (tmp[0] + tmp[1]) / 2, (tmp[2] + tmp[3]) / 2]
    return out

def main():
    old, new = load_old(), load_new()
    names = list(old) + list(new)
    M = np.array([{**old, **new}[n] for n in names], float)

    Z = (M - M.mean(0)) / M.std(0, ddof=1)
    comb = Z.mean(1)
    sem = Z.std(1, ddof=1) / np.sqrt(Z.shape[1])
    best = Z.argmax(0)

    print(f"{'Method':16s} " + " ".join(f"{c:>15s}" for c in COLS) + f" {'Avg':>7s} {'SEM':>6s}")
    for i, n in enumerate(names):
        star = "".join("*" if best[j] == i else " " for j in range(4))
        print(f"{n:16s} " + " ".join(f"{Z[i,j]:15.2f}" for j in range(4))
              + f" {comb[i]:7.2f} {sem[i]:6.2f}   best:{star}")

    ### LATEX ###
    disp = {"NB 10 uM": "NB 10 $\\mu$M", "RF (4096)": "RF (4096)", "DNN (4096)": "DNN (4096)"}
    lines = []
    lines.append("\\begin{table}[ht]")
    lines.append("\\centering")
    lines.append("\\caption{Overview of the performance of the benchmarked methods expressed as "
                 "z-scores per experiment, for the original Beyond-the-Hype models together with "
                 "the new 4096-feature experiments (RF, DNN and RandomNets). Each entry is the "
                 "combined (global and per-target) z-score, computed jointly across all methods per "
                 "column so that each column sums to zero. In italics the best performance per column "
                 "is highlighted.}")
    lines.append("\\label{tab:zscores}")
    lines.append("\\begin{tabular}{lcccccc}")
    lines.append("\\toprule")
    lines.append("Method & MCC & BEDROC & MCC & BEDROC & Average & SEM \\\\")
    lines.append(" & random & random & temporal & temporal & & \\\\")
    lines.append("\\midrule")
    for i, n in enumerate(names):
        if n == "RF (4096)":
            lines.append("\\midrule")
        label = disp.get(n, n).replace("_", "\\_")
        cells = []
        for j in range(4):
            # A z-score that rounds to zero from below prints as "-0.00", which
            # reads as a typo
            s = f"{Z[i,j]:.2f}".replace("-0.00", "0.00")
            cells.append(f"\\textit{{{s}}}" if best[j] == i else s)
        row = f"{label} & " + " & ".join(cells) + f" & {comb[i]:.2f} & {sem[i]:.2f} \\\\"
        lines.append(row)
    lines.append("\\bottomrule")
    lines.append("\\end{tabular}")
    lines.append("\\end{table}")
    tex = "\n".join(lines)
    with open("output/zscore_table.tex", "w") as fh:
        fh.write(tex + "\n")
    print("\nsaved output/zscore_table.tex")
